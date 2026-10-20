import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Palette
    COLOR_PRIMARY = RGBColor(24, 76, 120)     # Deep Royal Navy #184C78
    COLOR_SECONDARY = RGBColor(13, 148, 136)  # Teal #0D9488
    COLOR_TEXT = RGBColor(40, 44, 52)         # Slate Dark #282C34
    COLOR_MUTED = RGBColor(100, 116, 139)     # Slate Grey #64748B

    # Document Header / Badge
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_meta = p_meta.add_run("INAMIGOS FOUNDATION INTERNSHIP  |  TASK 1 SUBMISSION")
    run_meta.font.name = "Arial"
    run_meta.font.size = Pt(8.5)
    run_meta.font.bold = True
    run_meta.font.color.rgb = COLOR_MUTED

    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("AI Tools Discovery Report")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(16)
    run_sub = p_sub.add_run("Exploration of 5 Transformative AI Tools Across Functional Domains")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = COLOR_SECONDARY

    # Metadata Box
    meta_table = doc.add_table(rows=1, cols=3)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    
    widths = [Inches(2.3), Inches(2.3), Inches(2.3)]
    for i, col in enumerate(meta_table.columns):
        col.width = widths[i]

    row_cells = meta_table.rows[0].cells
    fields = [
        ("Candidate / Intern:", "Prajin K"),
        ("Program / Role:", "InAmigos Foundation Intern"),
        ("Domain:", "AI & Emerging Technologies")
    ]
    for idx, (label, val) in enumerate(fields):
        set_cell_background(row_cells[idx], "F1F5F9")
        set_cell_margins(row_cells[idx], top=120, bottom=120, left=140, right=140)
        p = row_cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p.add_run(f"{label}\n")
        r1.font.name = "Arial"
        r1.font.size = Pt(8.5)
        r1.font.color.rgb = COLOR_MUTED
        r2 = p.add_run(val)
        r2.font.name = "Arial"
        r2.font.size = Pt(10)
        r2.font.bold = True
        r2.font.color.rgb = COLOR_PRIMARY

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 1. Project Background & Objective
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("1. Executive Summary & Objective")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(15)
    r_h1.font.bold = True
    r_h1.font.color.rgb = COLOR_PRIMARY

    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.line_spacing = 1.15
    p_intro.paragraph_format.space_after = Pt(8)
    r_intro = p_intro.add_run(
        "Artificial Intelligence has rapidly evolved from speculative research into indispensable daily software. "
        "Modern professionals, researchers, and creators leverage specialized generative models to eliminate manual drudgery, "
        "elevate creative ideation, and dramatically accelerate execution speed. The objective of this discovery report is to "
        "evaluate five distinct AI tools spanning diverse functional categories—Education, Productivity, Image Generation, "
        "Video Creation, and Presentation Design—identifying their underlying purpose, practical workflow applications, and transformative impact."
    )
    r_intro.font.name = "Arial"
    r_intro.font.size = Pt(10.5)
    r_intro.font.color.rgb = COLOR_TEXT

    # Tools definition list
    tools_data = [
        {
            "num": "2.1",
            "name": "Perplexity AI",
            "category": "Education & Academic Research",
            "purpose": "An AI-powered conversational search engine and academic research copilot designed to provide verified, synthesized answers anchored in real-time web citations.",
            "use_case_title": "Interactive Academic Literature Synthesis",
            "use_case": (
                "A student researching 'The Economic and Grid Impacts of Distributed Solar Storage in Developing Nations' "
                "queries Perplexity. Instead of wading through millions of unranked search links, Perplexity synthesizes insights "
                "from recent whitepapers and policy journals, delivering concise paragraphs annotated with clickable numeric citations. "
                "The student can switch to 'Academic Mode' to restrict sources strictly to peer-reviewed repositories like arXiv and Semantic Scholar."
            ),
            "key_features": [
                "Direct Source Attribution: Every assertion is backed by clickable hyperlinks to primary sources.",
                "Collections & Thread Sharing: Organizes research projects into structured hubs.",
                "Multi-modal Research: Capable of analyzing user-uploaded PDFs, data sheets, and charts."
            ],
            "impact": "Reduces preliminary literature review time by over 70% while drastically mitigating factual hallucination."
        },
        {
            "num": "2.2",
            "name": "Notion AI",
            "category": "Productivity & Workspace Management",
            "purpose": "A connected workflow and knowledge-base AI assistant built natively into the Notion productivity workspace.",
            "use_case_title": "Meeting Minutes & Automated Action Item Tracking",
            "use_case": (
                "Following a 45-minute multi-stakeholder project sync, an intern pastes raw bullet points and transcript notes into a Notion page. "
                "By invoking Notion AI, the tool instantly generates an executive recap, creates a structured table of deliverables with deadlines, "
                "and tags assigned team members. Furthermore, users can ask Notion AI questions across their entire company workspace database."
            ),
            "key_features": [
                "Automated Synthesis: Summarizes extensive documentation, roadmaps, and requirements docs.",
                "Q&A Across Knowledge Bases: Connects with Slack, Google Drive, and internal Notion wikis.",
                "Writing & Tone Adaptation: Refactors informal drafts into professional proposals or client briefs."
            ],
            "impact": "Transforms disorganized documentation into actionable team workflows without switching apps."
        },
        {
            "num": "2.3",
            "name": "Midjourney",
            "category": "Image Generation & Visual Arts",
            "purpose": "A state-of-the-art diffusion-based text-to-image synthesis model engineered for artistic conceptualization, photorealism, and visual asset generation.",
            "use_case_title": "Rapid Campaign Concept Art & Mood Boarding",
            "use_case": (
                "A creative intern planning an environmental awareness campaign needs high-impact visual banners. "
                "By entering a descriptive prompt into Midjourney (e.g., '/imagine cinematic photograph of a child planting a glowing green sprout in a futuristic metropolis, golden hour volumetric lighting, shot on 35mm lens --ar 16:9 --v 6.1'), "
                "the engine outputs four high-resolution conceptual renders within 60 seconds, saving hours of stock photography scouring or costly studio shoots."
            ),
            "key_features": [
                "Unmatched Aesthetic Fidelity: Delivers industry-standard lighting, skin textures, and cinematic compositions.",
                "Granular Control Parameters: Fine-tunes aspect ratio (--ar), styling weight (--s), and camera focus.",
                "Inpainting & Outpainting: Enables selective redrawing ('Vary Region') and canvas expansion without quality degradation."
            ],
            "impact": "Democratizes high-end visual production, allowing non-designers to prototype creative assets rapidly."
        },
        {
            "num": "2.4",
            "name": "Synthesia",
            "category": "Video Creation & AI Avatars",
            "purpose": "An enterprise-grade AI video creation platform that converts written scripts into studio-quality videos featuring realistic digital avatars and expressive voice synthesis.",
            "use_case_title": "Multilingual Volunteer Onboarding & Training Modules",
            "use_case": (
                "An NGO or corporate coordinator needs to produce an onboarding series on 'Child Safety Protocols' for global volunteers across 12 countries. "
                "Rather than booking video production equipment, lighting, and actors, the team pastes script chapters into Synthesia. "
                "Synthesia renders an articulate photorealistic avatar presenting the lesson in Hindi, English, Spanish, and French with natural lip-syncing."
            ),
            "key_features": [
                "150+ Diverse AI Avatars: Fully modeled human avatars with natural micro-expressions.",
                "120+ Languages and Accents: Studio-quality text-to-speech with customizable pacing and pauses.",
                "Rapid Iteration: Script updates require editing text and re-rendering rather than re-shooting footage."
            ],
            "impact": "Cuts corporate video production budgets and timelines by over 80% while enabling seamless localization."
        },
        {
            "num": "2.5",
            "name": "Gamma AI (Gamma.app)",
            "category": "Presentations & Interactive Decks",
            "purpose": "An AI-native presentation, document, and webpage generator that constructs beautifully formatted, responsive visual slide decks from simple textual prompts or outlines.",
            "use_case_title": "Instant Pitch Deck & Educational Slide Generation",
            "use_case": (
                "An intern has 30 minutes to prepare a presentation summarizing 'Key Initiatives of InAmigos Foundation'. "
                "Inputting a rough markdown outline into Gamma AI generates a 10-card interactive slide deck equipped with consistent color palettes, "
                "smart visual cards, embedded charts, and icons. The slides adapt responsively to mobile, tablet, or projector displays."
            ),
            "key_features": [
                "One-Click Prompt-to-Deck: Builds complete narrative flows and layouts from single prompts or bullet notes.",
                "Fluid Card Architecture: Bypasses rigid slide boundaries with scrollable, interactive presentation blocks.",
                "Rich Media Integrations: Seamlessly embeds interactive widgets, videos, Figma files, and forms."
            ],
            "impact": "Eliminates tedious slide formatting, allowing presenters to focus purely on narrative substance."
        }
    ]

    h_tools = doc.add_heading(level=1)
    r_htools = h_tools.add_run("2. Detailed Exploration of 5 AI Tools")
    r_htools.font.name = "Arial"
    r_htools.font.size = Pt(15)
    r_htools.font.bold = True
    r_htools.font.color.rgb = COLOR_PRIMARY

    for tool in tools_data:
        h2 = doc.add_heading(level=2)
        h2.paragraph_format.space_before = Pt(10)
        h2.paragraph_format.space_after = Pt(2)
        r_h2 = h2.add_run(f"{tool['num']} {tool['name']}")
        r_h2.font.name = "Arial"
        r_h2.font.size = Pt(13)
        r_h2.font.bold = True
        r_h2.font.color.rgb = COLOR_PRIMARY

        # Category Badge
        p_badge = doc.add_paragraph()
        p_badge.paragraph_format.space_after = Pt(4)
        r_cat_label = p_badge.add_run("Functional Category: ")
        r_cat_label.font.name = "Arial"
        r_cat_label.font.size = Pt(9.5)
        r_cat_label.font.bold = True
        r_cat_label.font.color.rgb = COLOR_SECONDARY
        r_cat_val = p_badge.add_run(tool['category'])
        r_cat_val.font.name = "Arial"
        r_cat_val.font.size = Pt(9.5)
        r_cat_val.font.color.rgb = COLOR_TEXT

        # Purpose
        p_pur = doc.add_paragraph()
        p_pur.paragraph_format.space_after = Pt(4)
        r_pur_lbl = p_pur.add_run("Purpose: ")
        r_pur_lbl.font.name = "Arial"
        r_pur_lbl.font.size = Pt(10)
        r_pur_lbl.font.bold = True
        r_pur_lbl.font.color.rgb = COLOR_PRIMARY
        r_pur_txt = p_pur.add_run(tool['purpose'])
        r_pur_txt.font.name = "Arial"
        r_pur_txt.font.size = Pt(10)
        r_pur_txt.font.color.rgb = COLOR_TEXT

        # Callout Box for Practical Use Case
        use_table = doc.add_table(rows=1, cols=1)
        use_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        use_table.autofit = False
        use_table.columns[0].width = Inches(6.9)
        cell = use_table.rows[0].cells[0]
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

        cp = cell.paragraphs[0]
        cp.paragraph_format.space_after = Pt(2)
        cr_lbl = cp.add_run(f"Practical Real-World Use Case: {tool['use_case_title']}\n")
        cr_lbl.font.name = "Arial"
        cr_lbl.font.size = Pt(9.5)
        cr_lbl.font.bold = True
        cr_lbl.font.color.rgb = COLOR_SECONDARY

        cr_txt = cp.add_run(tool['use_case'])
        cr_txt.font.name = "Arial"
        cr_txt.font.size = Pt(9.5)
        cr_txt.font.color.rgb = COLOR_TEXT

        # Key Capabilities Bullet Points
        p_feat_lbl = doc.add_paragraph()
        p_feat_lbl.paragraph_format.space_before = Pt(4)
        p_feat_lbl.paragraph_format.space_after = Pt(2)
        r_flbl = p_feat_lbl.add_run("Key Features:")
        r_flbl.font.name = "Arial"
        r_flbl.font.size = Pt(9.5)
        r_flbl.font.bold = True
        r_flbl.font.color.rgb = COLOR_PRIMARY

        for feat in tool['key_features']:
            bp = doc.add_paragraph(style='List Bullet')
            bp.paragraph_format.space_after = Pt(2)
            r_feat = bp.add_run(feat)
            r_feat.font.name = "Arial"
            r_feat.font.size = Pt(9.5)
            r_feat.font.color.rgb = COLOR_TEXT

        # Value Impact
        p_imp = doc.add_paragraph()
        p_imp.paragraph_format.space_after = Pt(10)
        r_implbl = p_imp.add_run("Workflow Impact: ")
        r_implbl.font.name = "Arial"
        r_implbl.font.size = Pt(9.5)
        r_implbl.font.bold = True
        r_implbl.font.color.rgb = COLOR_SECONDARY
        r_imptxt = p_imp.add_run(tool['impact'])
        r_imptxt.font.name = "Arial"
        r_imptxt.font.size = Pt(9.5)
        r_imptxt.font.italic = True
        r_imptxt.font.color.rgb = COLOR_TEXT

    # 3. Comparative Matrix
    h_matrix = doc.add_heading(level=1)
    r_hm = h_matrix.add_run("3. Cross-Tool Comparison Matrix")
    r_hm.font.name = "Arial"
    r_hm.font.size = Pt(15)
    r_hm.font.bold = True
    r_hm.font.color.rgb = COLOR_PRIMARY

    table_comp = doc.add_table(rows=6, cols=4)
    table_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_comp.autofit = False

    col_widths = [Inches(1.4), Inches(1.8), Inches(2.2), Inches(1.5)]
    for row in table_comp.rows:
        for i, c_w in enumerate(col_widths):
            row.cells[i].width = c_w

    headers = ["Tool Name", "Domain", "Primary Strength", "Ideal User Profile"]
    for i, h_text in enumerate(headers):
        c = table_comp.rows[0].cells[i]
        set_cell_background(c, "184C78")
        set_cell_margins(c, top=100, bottom=100, left=100, right=100)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    comp_rows = [
        ("Perplexity AI", "Education & Research", "Real-time citation synthesis & academic filtering", "Students, Researchers, Analysts"),
        ("Notion AI", "Productivity & Ops", "Native workspace summarization & custom Q&A", "Project Managers, Teams, Interns"),
        ("Midjourney", "Image Generation", "Unrivaled photorealism & artistic composition", "Graphic Designers, Marketers, Artists"),
        ("Synthesia", "Video Creation", "Multilingual text-to-avatar studio video", "Educators, HR, Corporate Trainers"),
        ("Gamma AI", "Presentations", "Rapid outline-to-responsive slide generation", "Speakers, Founders, Students")
    ]

    for row_idx, data in enumerate(comp_rows, start=1):
        bg_color = "FFFFFF" if row_idx % 2 == 1 else "F8FAFC"
        for col_idx, text in enumerate(data):
            cell = table_comp.rows[row_idx].cells[col_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80, left=90, right=90)
            p = cell.paragraphs[0]
            if col_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(text)
                r.font.bold = True
                r.font.color.rgb = COLOR_PRIMARY
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(text)
                r.font.color.rgb = COLOR_TEXT
            r.font.name = "Arial"
            r.font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 4. Key Learnings & Strategic Takeaways
    h_learn = doc.add_heading(level=1)
    r_hl = h_learn.add_run("4. Key Learnings & Future Perspective")
    r_hl.font.name = "Arial"
    r_hl.font.size = Pt(15)
    r_hl.font.bold = True
    r_hl.font.color.rgb = COLOR_PRIMARY

    learnings = [
        ("Specialization Outperforms Generalization: ", "While broad LLMs provide generic answers, specialized tools (e.g., Perplexity for academic citations or Gamma for slide layouts) offer tailored user interfaces that drastically reduce friction."),
        ("Multi-Modal Synergies: ", "Modern workflows combine multiple AI tools—using Perplexity for research, Notion AI for outlining, Midjourney for custom graphics, Synthesia for narration, and Gamma for presentation."),
        ("Human Oversight Remains Paramount: ", "AI accelerates drafting and synthesis, but human critical evaluation is essential to verify facts, maintain brand tone, and ensure ethical authenticity.")
    ]

    for title_str, body_str in learnings:
        lp = doc.add_paragraph(style='List Bullet')
        lp.paragraph_format.space_after = Pt(4)
        rt = lp.add_run(title_str)
        rt.font.name = "Arial"
        rt.font.size = Pt(10)
        rt.font.bold = True
        rt.font.color.rgb = COLOR_PRIMARY
        rb = lp.add_run(body_str)
        rb.font.name = "Arial"
        rb.font.size = Pt(10)
        rb.font.color.rgb = COLOR_TEXT

    # 5. Conclusion & Verification Details
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    box_concl = doc.add_table(rows=1, cols=1)
    box_concl.alignment = WD_TABLE_ALIGNMENT.CENTER
    box_concl.columns[0].width = Inches(6.9)
    cc = box_concl.rows[0].cells[0]
    set_cell_background(cc, "E6FFFA")
    set_cell_margins(cc, top=120, bottom=120, left=150, right=150)
    p_con = cc.paragraphs[0]
    r_c_head = p_con.add_run("Internship Submission Acknowledgment\n")
    r_c_head.font.name = "Arial"
    r_c_head.font.size = Pt(10.5)
    r_c_head.font.bold = True
    r_c_head.font.color.rgb = COLOR_SECONDARY
    r_c_body = p_con.add_run(
        "This project report was prepared by Prajin K as part of Task 1 (AI Tools Discovery) for the InAmigos Foundation Internship. "
        "All associated presentation assets, markdown repositories, and deliverables are archived in the designated Google Drive and GitHub repository."
    )
    r_c_body.font.name = "Arial"
    r_c_body.font.size = Pt(9.5)
    r_c_body.font.color.rgb = COLOR_TEXT

    doc.save("Task1_AI_Tools_Discovery.docx")
    print("Task1_AI_Tools_Discovery.docx generated successfully!")

if __name__ == "__main__":
    create_document()
